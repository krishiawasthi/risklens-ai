from io import BytesIO

from docx import Document

from tests.sample_result import sample_result
from utils.report import build_report, report_filename
from utils.scoring import load_config


def all_text(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_report_has_all_sections():
    text = all_text(build_report(sample_result(), load_config()))
    for heading in ["Operational Risk Review: Example Manufacturing Corp", "1. Summary", "2. Scope and sources",
                    "3. Prioritised risk register", "4. Risk matrix", "5. Financial indicators",
                    "6. Not assessed and data gaps", "7. Method and scoring assumptions",
                    "Annex A: Detailed findings", "Annex B: XBRL source values"]:
        assert heading in text


def test_report_separates_observed_and_ai():
    text = all_text(build_report(sample_result(), load_config()))
    assert "Observed evidence (quotations verified word for word against the filing)" in text
    assert "AI interpretation (not verified)" in text
    assert "Observed (deterministic rule)" in text
    assert "Recommended follow-up questions (fixed for this rule)" in text


def test_report_quotes_are_verbatim_and_referenced():
    result = sample_result()
    text = all_text(build_report(result, load_config()))
    for finding in result.extraction.findings:
        for e in finding.evidence:
            assert e.quote in text
            assert f"paragraph {e.paragraph_id}" in text


def test_report_lists_gaps_and_rejections():
    text = all_text(build_report(sample_result(), load_config()))
    assert "high_leverage" in text  # not assessed: Liabilities not tagged
    assert "Fabricated finding" in text  # rejected by the evidence check


def test_report_summary_numbers_match_register():
    result = sample_result()
    text = all_text(build_report(result, load_config()))
    n = len(result.flags) + len(result.extraction.findings)
    assert f"identified {n} potential risks" in text


def test_report_has_no_em_dashes():
    text = all_text(build_report(sample_result(), load_config()))
    assert "—" not in text


def test_report_without_ai_or_financials():
    result = sample_result(with_ai=False)
    text = all_text(build_report(result, load_config()))
    assert "was not run for this report" in text
    empty = result.model_copy(update={"current_period": None, "prior_period": None, "metrics": {}, "flags": []})
    text = all_text(build_report(empty, load_config()))
    assert "No XBRL financial data was available" in text


def test_report_filename():
    assert report_filename(sample_result()) == "RiskLens_EXMC_10-K_2025-12-31.docx"
