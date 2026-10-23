"""20-F support: foreign companies, IFRS data and non-USD currencies.
All data is synthetic; no network calls."""
import json
from datetime import date
from io import BytesIO

import pytest
from docx import Document

from tests.sample_sec import (
    ACC_10K,
    ACC_20F,
    CIK,
    FOREIGN_CIK,
    IFRS_COMPANYFACTS,
    COMPANYFACTS,
    FakeSession,
    foreign_routes,
    sample_20f_html,
)
from tests.test_llm_extractor import FakeGenerator, risk
from utils.display import fmt_metric, money
from utils.filing_text import display_item, html_to_lines, split_sections
from utils.financials import compute_all
from utils.pipeline import run_analysis
from utils.report import build_report
from utils.scoring import load_config
from utils.sec_client import FormNotFiled, SECClient
from utils.xbrl import detect_basis, select_periods

UA = "Test Person test@example.com"


def sec():
    return SECClient(user_agent=UA, session=FakeSession(foreign_routes()), cache_dir=None, sleep=lambda s: None)


# ---------------------------------------------------------------- sections

def sections_20f(html=None, items=("3D", "15", "16K", "5")):
    lines = html_to_lines(html or sample_20f_html())
    return {s.item: s for s in split_sections(lines, "20-F", list(items))}


def test_risk_factors_found_inside_item_3_without_d_prefix():
    s = sections_20f()["3.D"]
    assert s.found
    texts = [p.text for p in s.paragraphs]
    assert texts[0].startswith("Risk factors relating to our operations")  # a sentence, not the heading
    assert any("specialty chemicals" in t for t in texts)
    assert not any("Summary of Risk Factors" in t for t in texts)
    assert not any("Capitalization" in t for t in texts)
    assert not any("Form 20-F | 12" in t for t in texts)  # page footer removed
    assert not any("precision components" in t for t in texts)  # Item 4 not included


def test_risk_factors_found_with_d_prefix():
    html = sample_20f_html().replace(b"<div>Risk Factors</div>", b"<div>D. Risk Factors</div>")
    assert sections_20f(html)["3.D"].found


def test_item_16k_found_and_ends_at_next_item():
    s = sections_20f()["16K"]
    assert s.found and len(s.paragraphs) == 2
    assert not any("Item 18" in p.text for p in s.paragraphs)


def test_item_15_stops_at_item_16a():
    texts = [p.text for p in sections_20f()["15"].paragraphs]
    assert len(texts) == 2
    assert not any("financial expert" in t for t in texts)


def test_item_5_found_with_plural_reviews():
    assert sections_20f()["5"].found


def test_missing_risk_subheading_is_reported():
    html = sample_20f_html().replace(b"<div>Risk Factors</div>", b"<div>Other matters</div>")
    s = sections_20f(html, items=("3D",))["3.D"]
    assert not s.found
    assert "sub-heading" in s.note


def test_display_item():
    assert display_item("20-F", "3D") == "3.D"
    assert display_item("20-F", "16K") == "16K"
    assert display_item("10-K", "1A") == "1A"


# ---------------------------------------------------------------- IFRS data

def test_basis_detected_as_ifrs_in_euros():
    assert detect_basis(IFRS_COMPANYFACTS, ACC_20F) == ("ifrs-full", "EUR")
    assert detect_basis(COMPANYFACTS, ACC_10K) == ("us-gaap", "USD")


def test_ifrs_values_selected_with_currency():
    current, prior, notes = select_periods(IFRS_COMPANYFACTS, ACC_20F, date(2025, 12, 31), "20-F")
    assert current.currency == "EUR" and current.taxonomy == "ifrs-full"
    assert current.value("Revenue") == 5_000_000_000
    assert prior.value("Revenue") == 4_500_000_000
    assert current.facts["Revenue"].unit == "EUR"
    assert any("IFRS in EUR" in n for n in notes)


def test_finance_costs_used_and_flagged():
    current, _, notes = select_periods(IFRS_COMPANYFACTS, ACC_20F, date(2025, 12, 31), "20-F")
    assert current.facts["InterestExpense"].source_tag == "FinanceCosts"
    assert any("FinanceCosts" in n and "understated" in n for n in notes)


def test_missing_ifrs_liabilities_not_invented():
    current, prior, _ = select_periods(IFRS_COMPANYFACTS, ACC_20F, date(2025, 12, 31), "20-F")
    m = compute_all(current, prior)
    assert m["liabilities_to_equity"].status == "missing_input"
    assert m["current_ratio"].value == pytest.approx(0.8)
    assert m["interest_coverage"].value == pytest.approx(6.0)


def test_money_shows_currency():
    assert money(5_000_000_000, "EUR") == "EUR 5.0bn"
    assert money(-30_000_000, "TWD") == "-TWD 30.0m"
    assert money(73_733_000_000) == "$73.7bn"
    assert fmt_metric("stockholders_equity", 3_000_000_000, "EUR") == "EUR 3.0bn"


# ---------------------------------------------------------------- filing lookup

def test_latest_20f_skips_amendment():
    f = sec().latest_filing(FOREIGN_CIK, "20-F")
    assert f.accession == ACC_20F
    assert str(f.period_end) == "2025-12-31"


def test_10k_request_for_foreign_company_offers_20f():
    with pytest.raises(FormNotFiled) as e:
        sec().latest_filing(FOREIGN_CIK, "10-K", company="Example Overseas Plc")
    assert e.value.alternative == "20-F"


def test_20f_request_for_us_company_offers_10k():
    with pytest.raises(FormNotFiled) as e:
        sec().latest_filing(CIK, "20-F")
    assert e.value.alternative == "10-K"


def test_10q_request_for_foreign_company_explains():
    with pytest.raises(FormNotFiled, match="no 10-Q") as e:
        sec().latest_filing(FOREIGN_CIK, "10-Q")
    assert e.value.alternative is None


def test_40f_filer_is_not_supported():
    with pytest.raises(FormNotFiled, match="40-F") as e:
        sec().latest_filing("0000666666", "10-K")
    assert e.value.alternative is None


# ---------------------------------------------------------------- end to end

AI_20F = json.dumps({"risks": [
    risk("Specialty chemical supplier concentration", "third_party",
         [("3D-0002", "We depend on a small number of suppliers for specialty chemicals")]),
    risk("Supplier portal intrusion", "cybersecurity",
         [("16K-0002", "In 2025 we detected an intrusion into a supplier portal")]),
]})


def test_10k_request_switches_to_20f_end_to_end():
    gen = FakeGenerator(AI_20F)
    result = run_analysis("exov", "10-K", sec(), load_config(), gen)
    assert result.filing.form == "20-F"
    assert result.filing.accession == ACC_20F
    assert any("files annual reports on Form 20-F" in w for w in result.warnings)
    assert result.current_period.currency == "EUR"
    assert {s.item for s in result.sections if s.found} == {"3.D", "15", "16K"}
    assert "current_ratio_low" in {f.rule_id for f in result.flags}
    assert [f.title for f in result.extraction.findings] == [
        "Specialty chemical supplier concentration", "Supplier portal intrusion"]
    assert result.extraction.findings[0].evidence[0].section == "Item 3.D. Risk Factors"


def test_report_shows_euros_for_20f():
    result = run_analysis("EXOV", "20-F", sec(), load_config(), FakeGenerator(AI_20F))
    doc = Document(BytesIO(build_report(result, load_config())))
    text = "\n".join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for row in t.rows for c in row.cells])
    assert "EUR 5.0bn" in text
    assert "ifrs-full:Revenue" in text
    assert "Item 3.D" in text


# ---------------------------------------------------------------- SEC has no figures for the latest filing
# Real case seen with TSMC's 2025 20-F: company facts held ifrs-full figures
# only up to the previous 20-F; the latest filing had share counts only.

from tests.sample_sec import ACC_20F_PRIOR, facts_url, ffact, eur
from utils.xbrl import latest_filing_with_figures


def lagging_facts():
    """IFRS figures exist only in last year's 20-F (FY2024 with FY2023
    comparatives). A 20-F/A also carries a figure and must be ignored."""
    def prior(start, end, val):
        return ffact(start, end, val, accn=ACC_20F_PRIOR, filed="2025-04-11")
    old = lambda s, e, v: {**prior(s, e, v), "fy": 2024}
    return {
        "facts": {
            "dei": {"EntityCommonStockSharesOutstanding": {"units": {"shares": [
                {"end": "2026-03-31", "val": 1000, "accn": ACC_20F, "form": "20-F", "filed": "2026-04-10"}]}}},
            "ifrs-full": {
                "Revenue": eur(old("2024-01-01", "2024-12-31", 4_500_000_000), old("2023-01-01", "2023-12-31", 4_000_000_000),
                               {**ffact("2025-01-01", "2025-12-31", 9, accn="0000777777-26-000007", filed="2026-05-02"),
                                "form": "20-F/A"}),
                "CurrentAssets": eur(old(None, "2024-12-31", 2_200_000_000), old(None, "2023-12-31", 2_100_000_000)),
                "CurrentLiabilities": eur(old(None, "2024-12-31", 2_000_000_000), old(None, "2023-12-31", 1_900_000_000)),
            },
        }
    }


def test_latest_filing_with_figures_finds_previous_20f():
    assert latest_filing_with_figures(lagging_facts(), "20-F", ACC_20F, date(2026, 4, 10)) == (
        ACC_20F_PRIOR, date(2024, 12, 31))


def test_latest_filing_with_figures_ignores_later_and_other_forms():
    assert latest_filing_with_figures(lagging_facts(), "20-F", ACC_20F, date(2025, 1, 1)) is None
    assert latest_filing_with_figures(lagging_facts(), "10-K", "x") is None


def test_no_figures_message_blames_sec_not_company():
    _, _, notes = select_periods(lagging_facts(), ACC_20F, date(2025, 12, 31), "20-F")
    assert "has no financial figures for filing" in notes[0]
    assert "publishes them later" in notes[0]


def lagging_sec():
    routes = foreign_routes()
    routes[facts_url(FOREIGN_CIK)] = [(200, lagging_facts())]
    return SECClient(user_agent=UA, session=FakeSession(routes), cache_dir=None, sleep=lambda s: None)


def test_pipeline_uses_previous_filing_and_labels_it():
    result = run_analysis("EXOV", "20-F", lagging_sec(), load_config(), None)
    assert result.filing.accession == ACC_20F  # text still from the latest filing
    assert result.figures_note and ACC_20F_PRIOR in result.figures_note and "2024-12-31" in result.figures_note
    assert result.figures_note in result.warnings
    assert result.current_period.label == "Year ended 2024-12-31"
    assert result.current_period.value("Revenue") == 4_500_000_000
    assert result.prior_period.value("Revenue") == 4_000_000_000
    assert all(f.accession == ACC_20F_PRIOR for f in result.current_period.facts.values())
    assert result.metrics["current_ratio"].value == pytest.approx(1.1)
    assert len(result.not_assessed) < 7


def test_pipeline_no_fallback_when_latest_filing_has_figures():
    result = run_analysis("EXOV", "20-F", sec(), load_config(), None)
    assert result.figures_note is None
    assert result.current_period.label == "Year ended 2025-12-31"


def test_pipeline_nothing_available_still_not_assessed():
    routes = foreign_routes()
    routes[facts_url(FOREIGN_CIK)] = [(200, {"facts": {"dei": {}}})]
    s = SECClient(user_agent=UA, session=FakeSession(routes), cache_dir=None, sleep=lambda s: None)
    result = run_analysis("EXOV", "20-F", s, load_config(), None)
    assert result.current_period is None and result.figures_note is None
    assert len(result.not_assessed) == 7


def test_report_and_app_show_the_note():
    result = run_analysis("EXOV", "20-F", lagging_sec(), load_config(), None)
    doc = Document(BytesIO(build_report(result, load_config())))
    text = "\n".join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for row in t.rows for c in row.cells])
    assert "Note on financial figures: SEC has not published" in text
    assert text.count(ACC_20F_PRIOR) >= 3
