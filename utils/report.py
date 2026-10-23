"""Consulting-style Word report built from a saved AnalysisResult.

Everything in the report comes from the result object: no extra AI call is
made. The summary text is assembled from counts and titles by code, so it
can only say what the register already says. Each section labels what is
observed (filing text, XBRL figures) and what is AI judgement.
"""
from __future__ import annotations

from io import BytesIO
from typing import List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from utils.display import CHANGE_METRICS, category_label, currency_of, fmt_metric, metric_label, money
from utils.filing_text import display_item
from utils.financials import compute_metrics
from utils.register import RegisterRow, build_register, matrix_cells, priority_counts
from utils.schemas import AnalysisResult
from utils.scoring import ScoringConfig

INK = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x52, 0x51, 0x4E)
HEADER_FILL = "E1E0D9"
# Teal tints, darkest for High. Light enough for black text; the priority is
# always written in the cell as well, so shading never carries meaning alone.
BAND_FILL = {"High": "99D5CD", "Medium": "C9E9E4", "Low": "EAF6F4"}


# ---------------------------------------------------------------- helpers

def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _table(doc, headers: List[str], rows: List[List[str]], widths_cm: Optional[List[float]] = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        _shade(cell, HEADER_FILL)
    for values in rows:
        cells = table.add_row().cells
        for i, v in enumerate(values):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(9)
    if widths_cm:
        _fix_widths(table, widths_cm)
    doc.add_paragraph()
    return table


def _fix_widths(table, widths_cm: List[float]) -> None:
    """Word and LibreOffice only honour column widths when the table layout is
    fixed and the width is set on both the grid columns and every cell."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for i, w in enumerate(widths_cm):
        table.columns[i].width = Cm(w)
        for cell in table.columns[i].cells:
            cell.width = Cm(w)


def _label(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.bold = True


def _muted(doc, text: str, italic: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.color.rgb = MUTED
    run.font.size = Pt(9)
    run.italic = italic
    return p


def _page_number_footer(section) -> None:
    p = section.footer.paragraphs[0]
    run = p.add_run("RiskLens prototype. Generated from public SEC data. Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


# ---------------------------------------------------------------- sections

def _summary(doc, result: AnalysisResult, rows: List[RegisterRow], config: ScoringConfig) -> None:
    f = result.filing
    counts = priority_counts(rows, config)
    n_rule = sum(r.basis == "Rule" for r in rows)
    n_ai = sum(r.basis == "AI" for r in rows)
    bands = ", ".join(f"{counts[b.name]} {b.name}" for b in config.priority_bands)

    doc.add_heading("1. Summary", level=1)
    text = (
        f"RiskLens identified {len(rows)} potential risks in the {f.form} filed by {f.company} "
        f"for the period ended {f.period_end}: {bands} priority. "
        f"{n_rule} come from deterministic rules applied to the company's XBRL financial data"
    )
    if result.extraction is not None:
        text += (
            f" and {n_ai} from an AI-assisted review of the filing text. Every AI finding is supported "
            "by quotations that were checked word for word against the filing."
        )
    else:
        text += ". The AI-assisted review of the filing text was not run for this report."
    doc.add_paragraph(text)
    if result.figures_note:
        p = doc.add_paragraph()
        _label(p, "Note on financial figures: ")
        p.add_run(result.figures_note)
    if result.not_assessed:
        doc.add_paragraph(
            f"{len(result.not_assessed)} financial rule(s) could not be assessed because the required "
            "figures were not reported. They are listed in section 6 and were not estimated."
        )
    if rows:
        p = doc.add_paragraph()
        _label(p, "Highest priorities")
        for r in rows[:5]:
            basis = "financial rule" if r.basis == "Rule" else f"AI finding, {r.confidence.lower()} confidence"
            doc.add_paragraph(
                f"{r.ref}  {r.title} ({r.priority}, score {r.score:.0f}, {basis})", style="List Number"
            )


def _scope(doc, result: AnalysisResult) -> None:
    f = result.filing
    doc.add_heading("2. Scope and sources", level=1)
    sections = "; ".join(
        f"Item {s.item} {s.title} ({s.chars:,} characters)" if s.found else f"Item {s.item} {s.title} (not found)"
        for s in result.sections
    )
    x = result.extraction
    rows = [
        ["Company", f"{f.company} (CIK {f.cik})"],
        ["Filing", f"{f.form}, period ended {f.period_end}, filed {f.filed}"],
        ["Accession number", f.accession],
        ["Document", f.url or "n/a"],
        ["Sections analysed", sections or "none"],
        ["Financial data", result.figures_note or "SEC XBRL company facts for the same filing"],
        ["AI model", ", ".join(x.models_used) if x and x.models_used else "not run"],
        ["Report generated", result.generated_at],
    ]
    _table(doc, ["Item", "Detail"], rows, [4.0, 13.0])


def _register(doc, rows: List[RegisterRow]) -> None:
    doc.add_heading("3. Prioritised risk register", level=1)
    _muted(doc, "Score = likelihood x impact x category weight. L = likelihood, I = impact (1 to 5). "
                "Rule rows are calculated from reported figures; AI rows are AI judgements supported by verified quotations.")
    if not rows:
        doc.add_paragraph("No risks were identified.")
        return
    table_rows = [
        [r.ref, r.title, category_label(r.category), r.basis, r.likelihood, r.impact, f"{r.score:.0f}", r.priority, r.confidence]
        for r in rows
    ]
    table = _table(doc, ["Ref", "Risk", "Category", "Basis", "L", "I", "Score", "Priority", "Confidence"],
                   table_rows, [1.3, 5.2, 2.6, 1.2, 0.7, 0.7, 1.4, 1.6, 2.3])
    for i, r in enumerate(rows, start=1):
        _shade(table.rows[i].cells[7], BAND_FILL.get(r.priority, "FFFFFF"))


def _matrix(doc, rows: List[RegisterRow], config: ScoringConfig) -> None:
    doc.add_heading("4. Risk matrix", level=1)
    cells = {(c["likelihood"], c["impact"]): c for c in matrix_cells(rows, config)}
    impacts = list(range(config.impact_scale.min, config.impact_scale.max + 1))
    likelihoods = list(range(config.likelihood_scale.max, config.likelihood_scale.min - 1, -1))
    table = doc.add_table(rows=len(likelihoods) + 1, cols=len(impacts) + 1)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Likelihood \\ Impact"
    for j, imp in enumerate(impacts, start=1):
        table.rows[0].cells[j].text = str(imp)
        _shade(table.rows[0].cells[j], HEADER_FILL)
    _shade(table.rows[0].cells[0], HEADER_FILL)
    for i, lik in enumerate(likelihoods, start=1):
        table.rows[i].cells[0].text = str(lik)
        _shade(table.rows[i].cells[0], HEADER_FILL)
        for j, imp in enumerate(impacts, start=1):
            c = cells[(lik, imp)]
            cell = table.rows[i].cells[j]
            cell.text = c["refs"]
            _shade(cell, BAND_FILL.get(c["band"], "FFFFFF"))
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8)
    _fix_widths(table, [3.2] + [2.76] * len(impacts))
    legend = ", ".join(f"{b.name} from score {b.min_score:g}" for b in config.priority_bands)
    _muted(doc, f"Cell shading shows the priority band for likelihood x impact at a category weight of 1.0 ({legend}).")


def _financials(doc, result: AnalysisResult) -> None:
    doc.add_heading("5. Financial indicators", level=1)
    if result.current_period is None:
        doc.add_paragraph("No XBRL financial data was available for this filing.")
        return
    if result.figures_note:
        doc.add_paragraph(result.figures_note)
    cur, pri = result.current_period, result.prior_period
    ccy = cur.currency
    prior_metrics = compute_metrics(pri) if pri else {}
    level_rows, change_rows = [], []
    for name, m in result.metrics.items():
        if name in CHANGE_METRICS:
            change_rows.append([metric_label(name), fmt_metric(name, m.value, ccy), m.note])
        else:
            p = prior_metrics.get(name)
            level_rows.append([metric_label(name), fmt_metric(name, m.value, ccy), fmt_metric(name, p.value if p else None, ccy), m.note])
    _table(doc, ["Indicator", cur.label, pri.label if pri else "Prior period", "Note"], level_rows, [5.5, 3.0, 3.0, 5.5])
    if change_rows:
        _table(doc, ["Change versus prior period", "Value", "Note"], change_rows, [6.0, 3.0, 8.0])
    _muted(doc, "All figures are calculated in Python from values tagged in the filing's XBRL data. "
                "Missing inputs are shown as n/a and are never estimated. Source tags are listed in Annex B.")


def _rule_detail(doc, r: RegisterRow, currency: str = "USD") -> None:
    f = r.flag
    p = doc.add_paragraph()
    _label(p, "Observed (deterministic rule): ")
    p.add_run(
        f"{metric_label(f.metric)} was {fmt_metric(f.metric, f.observed_value, currency)} for the {f.period.lower()}. "
        f"The rule '{f.rule_id}' fires when the value is {f.operator} {fmt_metric(f.metric, f.threshold, currency)}."
    )
    p = doc.add_paragraph()
    _label(p, "Source: ")
    p.add_run("XBRL data in accession " + ", ".join(f.sources) if f.sources else "XBRL data")
    p = doc.add_paragraph()
    _label(p, "Rating: ")
    p.add_run(f"likelihood {f.likelihood}, impact {f.impact}, set in config/scoring.yaml for this rule.")
    if f.follow_up_questions:
        p = doc.add_paragraph()
        _label(p, "Recommended follow-up questions (fixed for this rule)")
        for q in f.follow_up_questions:
            doc.add_paragraph(q, style="List Bullet")
    p = doc.add_paragraph()
    _label(p, "Limitations: ")
    p.add_run("A single ratio cannot show why it is at this level. The threshold is a configurable "
              "assumption, not an industry standard, and may not suit every business model.")


def _ai_detail(doc, r: RegisterRow) -> None:
    x = r.finding
    p = doc.add_paragraph()
    _label(p, "Observed evidence (quotations verified word for word against the filing)")
    for e in x.evidence:
        q = doc.add_paragraph(style="Quote")
        q.add_run(f"“{e.quote}”")
        _muted(doc, f"{e.section}, paragraph {e.paragraph_id}, {e.filing.form} accession {e.filing.accession}")
    p = doc.add_paragraph()
    _label(p, "AI interpretation (not verified): ")
    p.add_run(x.interpretation)
    p = doc.add_paragraph()
    _label(p, "AI rating: ")
    p.add_run(f"likelihood {x.likelihood}, impact {x.impact}. {x.rationale}")
    p = doc.add_paragraph()
    _label(p, f"Confidence {x.confidence}: ")
    reason = x.confidence_reason.rstrip(".") + "."
    if x.corroborated_by:
        reason += f" Related financial rule(s): {', '.join(x.corroborated_by)}."
    p.add_run(reason + " Confidence is computed from the verified evidence, not rated by the AI.")
    p = doc.add_paragraph()
    _label(p, "Recommended follow-up questions (AI-suggested)")
    for q in x.follow_up_questions:
        doc.add_paragraph(q, style="List Bullet")
    p = doc.add_paragraph()
    _label(p, "Limitations")
    for lim in x.limitations:
        doc.add_paragraph(lim, style="List Bullet")


def _details(doc, rows: List[RegisterRow], currency: str = "USD") -> None:
    doc.add_heading("Annex A: Detailed findings", level=1)
    for r in rows:
        doc.add_heading(f"{r.ref}  {r.title}", level=2)
        _muted(doc, f"{category_label(r.category)} | {r.priority} priority | score {r.score:.0f} | "
                    f"{'Financial rule' if r.basis == 'Rule' else 'AI finding'}")
        if r.basis == "Rule":
            _rule_detail(doc, r, currency)
        else:
            _ai_detail(doc, r)


def _gaps(doc, result: AnalysisResult) -> None:
    doc.add_heading("6. Not assessed and data gaps", level=1)
    items = []
    items += [f"Rule '{n.rule_id}' not assessed: {n.reason}." for n in result.not_assessed]
    items += [f"Item {s.item} ({s.title}) was not found in the document and was not analysed."
              for s in result.sections if not s.found]
    items += result.data_notes
    if result.extraction:
        items += [f"AI finding rejected: '{r.title}'. {r.reason}." for r in result.extraction.rejected]
        items += result.extraction.warnings
    items += result.warnings
    if not items:
        doc.add_paragraph("None recorded.")
    for item in items:
        doc.add_paragraph(item, style="List Bullet")
    p = doc.add_paragraph()
    _label(p, "General limitations")
    for text in [
        "The review uses public disclosures only. It has no access to internal controls, incident logs or testing evidence.",
        "Risk factor disclosures describe possible events and do not confirm that they occurred.",
        "AI interpretations, ratings and suggested questions require review by a qualified person.",
        "This report is a screening aid. It is not an audit opinion or investment advice.",
    ]:
        doc.add_paragraph(text, style="List Bullet")


def _method(doc, config: ScoringConfig, form: str) -> None:
    doc.add_heading("7. Method and scoring assumptions", level=1)
    c = config.confidence
    bands = "; ".join(f"{b.name}: score of {b.min_score:g} or more" for b in config.priority_bands)
    weights = ", ".join(f"{category_label(k)} {v:g}" for k, v in config.category_weights.items())
    for text in [
        "Score = likelihood x impact x category weight, on 1 to 5 scales.",
        f"Priority bands: {bands}.",
        f"Category weights: {weights}.",
        f"Confidence (AI findings): High with verified quotations from {c.high_min_paragraphs} or more paragraphs, "
        f"or {c.high_min_paragraphs_if_corroborated} or more when a financial rule in the same category also fired; "
        f"Medium with {c.medium_min_paragraphs}; Low with 1. Findings with no verified quotation are rejected.",
        f"A quotation must be at least {config.evidence.min_quote_words} words and appear word for word in the cited paragraph.",
        f"Sections sent to the AI for a {form}: "
        + ", ".join(f"Item {display_item(form, i)}" for i in config.extraction.sections.get(form, [])) + ".",
        "All thresholds and weights are prototype assumptions held in config/scoring.yaml, not industry standards.",
    ]:
        doc.add_paragraph(text, style="List Bullet")


def _appendix(doc, result: AnalysisResult) -> None:
    if result.current_period is None:
        return
    doc.add_heading("Annex B: XBRL source values", level=1)
    pri = result.prior_period
    rows = []
    for concept, fact in result.current_period.facts.items():
        prior = pri.facts.get(concept) if pri else None
        accns = [fact.accession] + ([prior.accession] if prior and prior.accession != fact.accession else [])
        rows.append([concept, f"{result.current_period.taxonomy}:{fact.source_tag or ''}",
                     money(fact.value, fact.unit), money(prior.value, prior.unit) if prior else "n/a",
                     ", ".join(a for a in accns if a)])
    _table(doc, ["Concept", "XBRL tag", "Current", "Prior", "Accession"], rows, [3.2, 6.0, 2.2, 2.2, 3.4])


# ---------------------------------------------------------------- entry point

def build_report(result: AnalysisResult, config: ScoringConfig) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)  # A4
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.0))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    _page_number_footer(section)

    f = result.filing
    doc.add_heading(f"Operational Risk Review: {f.company}", level=0)
    _muted(doc, f"{f.form} for the period ended {f.period_end} | Accession {f.accession} | Prepared by RiskLens AI (prototype)")
    _muted(doc, "This report was generated automatically from a public SEC filing. Financial figures and quotations "
                "are taken directly from the filing and its XBRL data. Interpretations, ratings and AI-suggested "
                "questions are AI-generated and must be reviewed by a qualified person.", italic=True)

    rows = build_register(result)
    _summary(doc, result, rows, config)
    _scope(doc, result)
    _register(doc, rows)
    _matrix(doc, rows, config)
    _financials(doc, result)
    _gaps(doc, result)
    _method(doc, config, f.form)
    _details(doc, rows, currency_of(result))
    _appendix(doc, result)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def report_filename(result: AnalysisResult) -> str:
    return f"RiskLens_{result.ticker}_{result.filing.form}_{result.filing.period_end}.docx"
